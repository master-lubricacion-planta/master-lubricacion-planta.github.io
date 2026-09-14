# -*- coding: utf-8 -*-
"""Parsea PROGRAMA CON_LUB_Y26W<nn>.xlsx -> plan_w<nn>.json (generaliza parse_w35.py).

Uso: python3 parse_plan.py 37 [ruta.xlsx]
     (por defecto busca ../descargas/PROGRAMA_CON_LUB_Y26W<nn>.xlsx)

Primera pasada de mapeo OT->pauta: match exacto (tag+desc normalizados) contra
TODAS las semanas ya verificadas en Maximo (plan_w*.json existentes). El resto
queda null y se lista para verificar en Maximo (apply_maximo.py).
Fechas: jueves a miercoles; ancla jueves 13-08-2026 = semana 33.
"""
import openpyxl, json, re, sys, os, glob, unicodedata, datetime
from collections import defaultdict

base = os.path.dirname(os.path.abspath(__file__))
if len(sys.argv) < 2 or not sys.argv[1].isdigit():
    sys.exit('uso: python3 parse_plan.py <semana> [archivo.xlsx]')
W = int(sys.argv[1])
SRC = sys.argv[2] if len(sys.argv) > 2 else os.path.join(base, '..', 'descargas', f'PROGRAMA_CON_LUB_Y26W{W}.xlsx')
DIAS = ['Jueves', 'Viernes', 'Sábado', 'Domingo', 'Lunes', 'Martes', 'Miércoles']
ANCLA = datetime.date(2026, 8, 13)  # jueves semana 33

def norm(s):
    s = str(s or '').replace('\xa0', ' ')  # el Excel usa espacios duros
    s = unicodedata.normalize('NFD', s).encode('ascii', 'ignore').decode()
    return re.sub(r'[^A-Z0-9]+', ' ', s.upper()).strip()

wb = openpyxl.load_workbook(SRC, data_only=True)
hoja = [s for s in wb.sheetnames if f'W{W}' in s.upper() and 'LUB' in s.upper()]
if not hoja: sys.exit('no encuentro la hoja CON_LUB_Y26W%d en %s' % (W, wb.sheetnames))
ws = wb[hoja[0]]

ini = ANCLA + datetime.timedelta(days=(W - 33) * 7)
fechas = {(ini + datetime.timedelta(days=i)): DIAS[i] for i in range(7)}
days = [{'nombre': DIAS[i], 'fecha': (ini + datetime.timedelta(days=i)).strftime('%d.%m'), 'ots': []} for i in range(7)]
idx = {DIAS[i]: i for i in range(7)}

# columnas ancladas por encabezado (fila 5), nunca por posicion fija
hdr = {norm(ws.cell(row=5, column=c).value): c for c in range(1, 15) if ws.cell(row=5, column=c).value}
def col(*nombres):
    for n in nombres:
        for k, c in hdr.items():
            if norm(n) in k: return c
    sys.exit('falta columna %s en fila 5: %s' % (nombres, list(hdr)))
C_OT, C_ESP, C_DESC, C_TAG, C_INI = col('Orden de trabajo'), col('Especialidad'), col('Descripci'), col('Ubicaci'), col('Inicio')

fuera_rango = []; vistos = set()
for r in range(6, ws.max_row + 1):
    ot = ws.cell(row=r, column=C_OT).value
    if ot is None: continue
    ot = str(ot).strip()
    if not ot.isdigit(): continue
    esp = str(ws.cell(row=r, column=C_ESP).value or '').strip()
    desc = str(ws.cell(row=r, column=C_DESC).value or '').replace('\xa0', ' ').strip()
    tag = str(ws.cell(row=r, column=C_TAG).value or '').strip()
    hini = ws.cell(row=r, column=C_INI).value
    fecha = hini.date() if hasattr(hini, 'date') else None
    if fecha not in fechas:
        fuera_rango.append((ot, str(hini))); continue
    if ot in vistos: continue
    vistos.add(ot)
    days[idx[fechas[fecha]]]['ots'].append({'ot': ot, 'tag': tag, 'desc': desc, 'esp': esp, 'pauta': None})

# --- mapeo primera pasada desde las semanas ya verificadas
por_tagdesc = {}
for f in sorted(glob.glob(os.path.join(base, 'plan_w*.json'))):
    if f.endswith(f'plan_w{W}.json'): continue
    for d in json.load(open(f, encoding='utf-8'))['days']:
        for o in d['ots']:
            if o.get('pauta'):
                por_tagdesc[(norm(o['tag']), norm(o['desc']))] = o['pauta']

map_td = 0; pend = []
for d in days:
    for o in d['ots']:
        if o['esp'] != 'LUB':
            continue  # otra especialidad: sin pauta digital
        k = (norm(o['tag']), norm(o['desc']))
        if k in por_tagdesc:
            o['pauta'] = por_tagdesc[k]; map_td += 1
        else:
            pend.append((d['nombre'], o['ot'], o['tag'], o['desc'][:45]))

plan = {'week': W, 'turno': 'A' if W % 2 == 1 else 'B', 'anio': 2026, 'days': days}
out = os.path.join(base, f'plan_w{W}.json')
json.dump(plan, open(out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

tot = sum(len(d['ots']) for d in days)
lub = sum(1 for d in days for o in d['ots'] if o['esp'] == 'LUB')
print(f'semana {W} ({days[0]["fecha"]} - {days[6]["fecha"]}): OTs {tot} (LUB {lub}) | mapeadas tag+desc: {map_td} | PENDIENTES Maximo: {len(pend)}')
print('por dia:', [(d['nombre'], len(d['ots']), sum(1 for o in d['ots'] if o['pauta'])) for d in days])
if fuera_rango: print('fuera de rango:', fuera_rango[:8])
print('\nPendientes para Maximo (OT):', ','.join(p[1] for p in pend))
for p in pend: print(' ', p)
